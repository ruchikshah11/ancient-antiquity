"""Admin user list exports: CSV, Excel and PDF (same columns, same filters as the Users page)."""
import csv
import io
import os

from django.http import HttpResponse
from django.utils import timezone

HEADERS = [
    "Name", "Email", "Phone", "Role", "Status", "Two-step login", "Bookings", "Assigned places", "Joined", "Last login",
]
# Relative PDF column widths (landscape A4).
PDF_WIDTHS = [1.3, 2.0, 1.05, 0.85, 0.6, 0.65, 0.7, 1.85, 0.85, 1.0]


def export_rows(users):
    """One row per user; dates as naive local datetimes so Excel shows them as dates."""
    for u in users:
        if u.role == u.Role.STAFF:
            places = ", ".join(m.name for m in u.assigned_museums.all())
        elif u.role == u.Role.PARTNER:
            places = ", ".join(m.name for m in u.partner_places.all())
        else:
            places = ""
        last = timezone.localtime(u.last_login).replace(tzinfo=None) if u.last_login else None
        yield [
            u.get_full_name() or "-",
            u.email,
            u.phone or "",
            u.get_role_display(),
            "Active" if u.is_active else "Disabled",
            "On" if u.totp_enabled else "Off",
            getattr(u, "booking_count", 0),
            places,
            timezone.localtime(u.date_joined).replace(tzinfo=None),
            last,
        ]


def _filename(ext):
    return f"users-{timezone.localdate():%Y-%m-%d}.{ext}"


def _describe(filters):
    parts = [f"{k}: {v}" for k, v in filters.items() if v]
    return ", ".join(parts) or "All users"


def as_csv(users):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename("csv")}"'
    response.write("\ufeff")  # BOM so Excel opens Hindi/other scripts as UTF-8
    writer = csv.writer(response)
    writer.writerow(HEADERS)
    for row in export_rows(users):
        row[8] = row[8].strftime("%Y-%m-%d")
        row[9] = row[9].strftime("%Y-%m-%d %H:%M") if row[9] else ""
        writer.writerow(row)
    return response


def as_xlsx(users, filters):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCEFFB")
    for row in export_rows(users):
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 9).number_format = "yyyy-mm-dd"
        ws.cell(r, 10).number_format = "yyyy-mm-dd hh:mm"
    for i, width in enumerate([24, 32, 16, 10, 10, 13, 10, 40, 12, 17], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("About")
    info.append(["Ancient Antiquity user export"])
    info.append(["Filters", _describe(filters)])
    info.append(["Users", ws.max_row - 1])
    info.append(["Generated", timezone.localtime().replace(tzinfo=None)])
    info.append(["Note", "Contains personal data. Store and share it only as your privacy policy allows."])
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 12
    info.column_dimensions["B"].width = 60

    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{_filename("xlsx")}"'
    return response


def _pdf_font():
    """A TTF that covers Indian scripts when available (Windows: Nirmala UI); otherwise Helvetica."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        # Windows 11 ships Nirmala UI as a font collection (.ttc); older Windows as separate .ttf files.
        ("Nirmala", r"C:\Windows\Fonts\Nirmala.ttc", r"C:\Windows\Fonts\Nirmala.ttc"),
        ("Nirmala", r"C:\Windows\Fonts\Nirmala.ttf", r"C:\Windows\Fonts\NirmalaB.ttf"),
        ("NotoSans", "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf", "/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf"),
        ("DejaVuSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ]
    for name, regular, bold in candidates:
        if os.path.exists(regular):
            if name not in pdfmetrics.getRegisteredFontNames():
                # In Nirmala.ttc, face 0 is Regular and face 1 is Bold.
                bold_index = 1 if bold == regular and regular.endswith(".ttc") else 0
                pdfmetrics.registerFont(TTFont(name, regular, subfontIndex=0))
                pdfmetrics.registerFont(TTFont(f"{name}-Bold", bold if os.path.exists(bold) else regular, subfontIndex=bold_index))
            return name, f"{name}-Bold"
    return "Helvetica", "Helvetica-Bold"


def as_pdf(users, filters):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    font, bold = _pdf_font()
    navy, sky = colors.HexColor("#0c4a6e"), colors.HexColor("#dcefFB")
    cell = ParagraphStyle("cell", fontName=font, fontSize=7.5, leading=9.5)
    head = ParagraphStyle("head", parent=cell, fontName=bold, textColor=navy)
    title = ParagraphStyle("title", fontName=bold, fontSize=15, leading=19, textColor=navy)
    sub = ParagraphStyle("sub", fontName=font, fontSize=8.5, leading=11, textColor=colors.HexColor("#5b6b82"))

    rows = list(export_rows(users))
    data = [[Paragraph(h, head) for h in HEADERS]]
    for r in rows:
        r[8] = r[8].strftime("%d %b %Y")
        r[9] = r[9].strftime("%d %b %Y %H:%M") if r[9] else "-"
        data.append([Paragraph(str(v).replace("&", "&amp;").replace("<", "&lt;"), cell) for v in r])

    page = landscape(A4)
    margin = 12 * mm
    usable = page[0] - 2 * margin
    widths = [usable * w / sum(PDF_WIDTHS) for w in PDF_WIDTHS]
    table = LongTable(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), sky),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, navy),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f9fd")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("ALIGN", (6, 1), (6, -1), "RIGHT"),
    ]))

    generated = timezone.localtime().strftime("%d %b %Y, %H:%M")

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 7.5)
        canvas.setFillColor(colors.HexColor("#5b6b82"))
        canvas.drawString(margin, 7 * mm, f"Ancient Antiquity · Users · generated {generated} · contains personal data")
        canvas.drawRightString(page[0] - margin, 7 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=page, leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=14 * mm,
                            title="Ancient Antiquity users", author="Ancient Antiquity")
    doc.build(
        [Paragraph("Users", title), Paragraph(f"{len(rows)} users · {_describe(filters)}", sub), Spacer(1, 4 * mm), table],
        onFirstPage=footer, onLaterPages=footer,
    )
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{_filename("pdf")}"'
    return response
