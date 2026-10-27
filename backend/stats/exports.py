"""Background export jobs: bookings (CSV / Excel / PDF) and users (CSV / Excel / PDF).

A job runs in a daemon thread inside the web server process (started after the creating request commits);
`manage.py run_export_jobs` picks up anything left pending or abandoned (e.g. the server restarted mid-run).
Files live in settings.EXPORTS_DIR (private: only the authenticated download endpoint serves them) and are
deleted after EXPORT_KEEP_DAYS.

Bookings are read in keyset chunks by id (no long-lived cursor, so progress updates never wait on SQLite),
written row by row: CSV streams straight to disk, Excel uses openpyxl write_only mode (a new sheet every
EXPORT_XLSX_SHEET_ROWS rows), PDF is capped at EXPORT_PDF_MAX_ROWS rows. User exports call the existing
accounts.exports functions.
"""
import csv
import datetime
import logging
import os
import threading
import traceback
import uuid
from pathlib import Path

from django.conf import settings
from django.db import close_old_connections, connection, transaction
from django.db.models import Q
from django.utils import timezone
from django.utils.dateparse import parse_date
from rest_framework.exceptions import PermissionDenied, ValidationError

from ancient_antiquity.dates import day_range
from bookings.models import Booking
from museums.models import Museum

from .models import ExportJob

log = logging.getLogger(__name__)

CHUNK = 2000
STALE_AFTER = datetime.timedelta(minutes=10)  # a running job with no heartbeat for this long was abandoned
MAX_ACTIVE_PER_USER = 3
BOOKING_KINDS = {ExportJob.Kind.BOOKINGS_CSV, ExportJob.Kind.BOOKINGS_XLSX, ExportJob.Kind.BOOKINGS_PDF}
USER_KINDS = {ExportJob.Kind.USERS_CSV, ExportJob.Kind.USERS_XLSX, ExportJob.Kind.USERS_PDF}
EXT = {"csv": "csv", "xlsx": "xlsx", "pdf": "pdf"}
BOOKING_FILTERS = ("from", "to", "museum", "search", "status", "channel", "is_group", "visit_date")
USER_FILTERS = ("search", "role", "is_active", "is_sample")


def exports_dir():
    return Path(settings.EXPORTS_DIR)


def _ext(kind):
    return kind.split("_", 1)[1]


# ------------------------------------------------------------------ filters -> querysets

def clean_filters(user, kind, raw):
    """Validate what the client sent; returns the filters to store (strings only, empty ones dropped)."""
    raw = raw if isinstance(raw, dict) else {}
    allowed = BOOKING_FILTERS if kind in BOOKING_KINDS else USER_FILTERS
    filters = {k: str(raw[k]).strip() for k in allowed if raw.get(k) not in (None, "")}
    if kind in BOOKING_KINDS:
        for key in ("from", "to", "visit_date"):
            if key in filters and not parse_date(filters[key]):
                raise ValidationError({key: "Use a date like 2026-10-25."})
        if "from" in filters and "to" in filters and filters["from"] > filters["to"]:
            raise ValidationError({"from": "The start date must be on or before the end date."})
        if "status" in filters and filters["status"] not in Booking.Status.values:
            raise ValidationError({"status": "Unknown status."})
        if "channel" in filters and filters["channel"] not in Booking.Channel.values:
            raise ValidationError({"channel": "Unknown channel."})
        bookings_queryset(user, filters)  # access checks (unknown place, place not assigned)
    else:
        if not user.is_admin_role:
            raise PermissionDenied("Only administrators can export users.")
    return filters


def _bool(value):
    return str(value).lower() in ("1", "true", "yes")


def bookings_queryset(user, filters):
    from bookings.views import ensure_museum_access, panel_bookings, search_bookings

    qs = panel_bookings(user)
    if filters.get("from") or filters.get("to"):
        first = parse_date(filters.get("from") or "") or datetime.date(2000, 1, 1)
        last = parse_date(filters.get("to") or "") or timezone.localdate()
        start, end = day_range(first, last)
        qs = qs.filter(created_at__gte=start, created_at__lt=end)
    if filters.get("museum"):
        museum = Museum.objects.filter(slug=filters["museum"]).first()
        if not museum:
            raise ValidationError({"museum": "Unknown place."})
        ensure_museum_access(user, museum)
        qs = qs.filter(museum=museum)
    if filters.get("status"):
        qs = qs.filter(status=filters["status"])
    if filters.get("channel"):
        qs = qs.filter(channel=filters["channel"])
    if filters.get("is_group"):
        qs = qs.filter(is_group=_bool(filters["is_group"]))
    if filters.get("visit_date"):
        qs = qs.filter(visit_date=parse_date(filters["visit_date"]))
    return search_bookings(qs, filters.get("search"))


def users_queryset(filters):
    from accounts.views import AdminUserViewSet

    qs = AdminUserViewSet.queryset.all()
    if filters.get("role"):
        qs = qs.filter(role=filters["role"])
    for key in ("is_active", "is_sample"):
        if filters.get(key):
            qs = qs.filter(**{key: _bool(filters[key])})
    for word in (filters.get("search") or "").split():  # like DRF SearchFilter on the Users page
        qs = qs.filter(Q(email__icontains=word) | Q(first_name__icontains=word) | Q(last_name__icontains=word)
                       | Q(phone__icontains=word))
    return qs.prefetch_related("assigned_museums")


def describe(job):
    f = job.filters
    if job.kind in USER_KINDS:
        return {"Search": f.get("search", ""), "Role": f.get("role", ""), "Active": f.get("is_active", "")}
    labels = {"from": "Booked from", "to": "Booked to", "museum": "Place", "search": "Search", "status": "Status",
              "channel": "Channel", "is_group": "Groups only", "visit_date": "Visit date"}
    return {labels[k]: v for k, v in f.items() if k in labels}


def file_name(job):
    stamp = timezone.localtime(job.created_at).strftime("%Y%m%d-%H%M")
    what = "users" if job.kind in USER_KINDS else "bookings"
    place = job.filters.get("museum")
    return f"{what}{'-' + place if place else ''}-{stamp}.{_ext(job.kind)}"


# ------------------------------------------------------------------ writers

def _heartbeat(job, rows):
    job.row_count = rows
    ExportJob.objects.filter(pk=job.pk).update(row_count=rows, heartbeat_at=timezone.now())


def iter_bookings(qs, job=None, limit=None):
    """Bookings in id order, CHUNK at a time (keyset), each chunk a short query."""
    from bookings.reports import EXPORT_RELATED

    last_id, done = 0, 0
    base = qs.select_related(*EXPORT_RELATED).order_by("id")
    while True:
        size = CHUNK if limit is None else min(CHUNK, limit - done)
        if size <= 0:
            return
        chunk = list(base.filter(id__gt=last_id)[:size])
        if not chunk:
            return
        for booking in chunk:
            yield booking
        done += len(chunk)
        last_id = chunk[-1].id
        if job is not None:
            _heartbeat(job, done)


def write_bookings_csv(job, qs, path):
    from bookings.reports import EXPORT_HEADERS, booking_row

    n = 0
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:  # BOM: Excel opens ₹/Hindi as UTF-8
        writer = csv.writer(fh)
        writer.writerow(EXPORT_HEADERS)
        for b in iter_bookings(qs, job):
            row = booking_row(b)
            row[1] = row[1].strftime("%Y-%m-%d %H:%M")
            writer.writerow(row)
            n += 1
    return n


def write_bookings_xlsx(job, qs, path):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    from bookings.reports import EXPORT_HEADERS, booking_row

    wb = Workbook(write_only=True)
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DCEFFB")
    widths = [13, 17, 12, 12, 34, 16, 18, 14, 18, 26, 9, 9, 11, 10, 15, 22, 13, 10, 14]
    per_sheet = settings.EXPORT_XLSX_SHEET_ROWS

    def new_sheet(number):
        ws = wb.create_sheet("Bookings" if number == 1 else f"Bookings {number}")
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        header = []
        for h in EXPORT_HEADERS:
            cell = WriteOnlyCell(ws, value=h)
            cell.font, cell.fill = bold, fill
            header.append(cell)
        ws.append(header)
        return ws

    sheets, n = 1, 0
    ws = new_sheet(1)
    for b in iter_bookings(qs, job):
        if n and n % per_sheet == 0:
            sheets += 1
            ws = new_sheet(sheets)
        ws.append(booking_row(b))
        n += 1

    info = wb.create_sheet("About")
    for row in (
        ["Ancient Antiquity bookings export"],
        *[[k, v] for k, v in describe(job).items()],
        ["Bookings", n],
        ["Generated", timezone.localtime().replace(tzinfo=None)],
    ):
        info.append(row)
    if sheets > 1:
        job.note = f"{n:,} rows over {sheets} sheets (Excel allows about 10 lakh rows per sheet)."
    wb.save(path)
    return n


def write_bookings_pdf(job, qs, path, total):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    from accounts.exports import _pdf_font
    from bookings.reports import EXPORT_HEADERS, booking_row

    cap = settings.EXPORT_PDF_MAX_ROWS
    columns = ["Reference", "Created", "Visit date", "Slot", "Place", "City", "Channel", "Visitors", "Amount",
               "Status", "Refund"]
    index = [EXPORT_HEADERS.index(c) for c in columns]
    rel = [1.0, 1.05, 0.8, 0.9, 2.2, 1.0, 0.9, 0.6, 0.8, 0.8, 0.7]

    font, bold = _pdf_font()
    navy, sky = colors.HexColor("#0c4a6e"), colors.HexColor("#dcefFB")
    cell = ParagraphStyle("cell", fontName=font, fontSize=7, leading=9)
    head = ParagraphStyle("head", parent=cell, fontName=bold, textColor=navy)
    title = ParagraphStyle("title", fontName=bold, fontSize=15, leading=19, textColor=navy)
    sub = ParagraphStyle("sub", fontName=font, fontSize=8.5, leading=11, textColor=colors.HexColor("#5b6b82"))

    def text(v):
        if isinstance(v, datetime.datetime):
            v = v.strftime("%d %b %Y %H:%M")
        elif isinstance(v, datetime.date):
            v = v.strftime("%d %b %Y")
        return Paragraph(str(v).replace("&", "&amp;").replace("<", "&lt;"), cell)

    data = [[Paragraph(h, head) for h in columns]]
    n = 0
    for b in iter_bookings(qs, job, limit=cap):
        row = booking_row(b)
        data.append([text(row[i]) for i in index])
        n += 1

    filters = ", ".join(f"{k}: {v}" for k, v in describe(job).items()) or "All bookings you can access"
    summary = f"{n:,} bookings · {filters}"
    if total > n:
        job.note = (f"PDF shows the first {n:,} of {total:,} bookings (PDF is limited to {cap:,} rows). "
                    "Use CSV or Excel for everything.")
        summary += f" · first {n:,} of {total:,} only: use CSV or Excel for all rows"

    page = landscape(A4)
    margin = 12 * mm
    usable = page[0] - 2 * margin
    table = LongTable(data, colWidths=[usable * w / sum(rel) for w in rel], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), sky),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, navy),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f9fd")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    generated = timezone.localtime().strftime("%d %b %Y, %H:%M")

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 7.5)
        canvas.setFillColor(colors.HexColor("#5b6b82"))
        canvas.drawString(margin, 7 * mm, f"Ancient Antiquity · Bookings · generated {generated}")
        canvas.drawRightString(page[0] - margin, 7 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=page, leftMargin=margin, rightMargin=margin, topMargin=margin,
                            bottomMargin=14 * mm, title="Ancient Antiquity bookings", author="Ancient Antiquity")
    doc.build([Paragraph("Bookings", title), Paragraph(summary, sub), Spacer(1, 4 * mm), table],
              onFirstPage=footer, onLaterPages=footer)
    return n


def write_users(job, path, total):
    from accounts import exports as user_exports

    users = users_queryset(job.filters)
    kind = _ext(job.kind)
    filters = describe(job)
    if kind == "csv":
        response = user_exports.as_csv(users)
    elif kind == "xlsx":
        response = user_exports.as_xlsx(users, filters)
    else:
        cap = settings.EXPORT_PDF_MAX_ROWS
        if total > cap:
            users = users[:cap]
            job.note = (f"PDF shows the first {cap:,} of {total:,} users (PDF is limited to {cap:,} rows). "
                        "Use CSV or Excel for everything.")
        response = user_exports.as_pdf(users, filters)
    with open(path, "wb") as fh:
        fh.write(response.content)
    return min(total, settings.EXPORT_PDF_MAX_ROWS) if kind == "pdf" else total


# ------------------------------------------------------------------ running

def claim(job_id):
    """Atomically move a pending job to running; None when another worker already has it."""
    now = timezone.now()
    taken = ExportJob.objects.filter(pk=job_id, status=ExportJob.Status.PENDING).update(
        status=ExportJob.Status.RUNNING, started_at=now, heartbeat_at=now, row_count=0, error="", note="",
    )
    return ExportJob.objects.select_related("user").get(pk=job_id) if taken else None


def run_job(job_id):
    job = claim(job_id)
    if job is None:
        return None
    folder = exports_dir() / str(job.user_id)
    folder.mkdir(parents=True, exist_ok=True)
    rel = f"{job.user_id}/{job.pk}-{uuid.uuid4().hex}.{_ext(job.kind)}"
    path = exports_dir() / rel
    partial = path.with_suffix(path.suffix + ".part")
    try:
        if job.kind in USER_KINDS:
            total = users_queryset(job.filters).count()
            ExportJob.objects.filter(pk=job.pk).update(total_rows=total)
            rows = write_users(job, partial, total)
        else:
            qs = bookings_queryset(job.user, job.filters)
            total = qs.count()
            ExportJob.objects.filter(pk=job.pk).update(total_rows=total)
            if job.kind == ExportJob.Kind.BOOKINGS_CSV:
                rows = write_bookings_csv(job, qs, partial)
            elif job.kind == ExportJob.Kind.BOOKINGS_XLSX:
                rows = write_bookings_xlsx(job, qs, partial)
            else:
                rows = write_bookings_pdf(job, qs, partial, total)
        os.replace(partial, path)
        ExportJob.objects.filter(pk=job.pk).update(
            status=ExportJob.Status.DONE, row_count=rows, file_path=rel, file_name=file_name(job),
            file_size=path.stat().st_size, note=job.note[:300], finished_at=timezone.now(), heartbeat_at=timezone.now(),
        )
    except Exception as exc:  # recorded on the job for the Exports page
        log.exception("Export job %s failed", job.pk)
        partial.unlink(missing_ok=True)
        detail = exc.detail if hasattr(exc, "detail") else str(exc)
        ExportJob.objects.filter(pk=job.pk).update(
            status=ExportJob.Status.FAILED, finished_at=timezone.now(),
            error=f"{detail}"[:2000] or traceback.format_exc()[-2000:],
        )
    return ExportJob.objects.get(pk=job.pk)


def _thread_main(job_id):
    try:
        run_job(job_id)
    finally:
        connection.close()


def start(job_id):
    """Run a job now: inline in tests (EXPORT_RUN_INLINE), else in a daemon thread after the commit."""
    if settings.EXPORT_RUN_INLINE:
        run_job(job_id)
        return

    def launch():
        threading.Thread(target=_thread_main, args=(job_id,), name=f"export-{job_id}", daemon=True).start()

    transaction.on_commit(launch)


def requeue_stale():
    """Running jobs without a heartbeat for STALE_AFTER were abandoned (server restarted): run them again."""
    return ExportJob.objects.filter(
        status=ExportJob.Status.RUNNING, heartbeat_at__lt=timezone.now() - STALE_AFTER,
    ).update(status=ExportJob.Status.PENDING)


def cleanup_expired():
    """Delete files older than EXPORT_KEEP_DAYS; the job stays listed as expired."""
    cutoff = timezone.now() - datetime.timedelta(days=settings.EXPORT_KEEP_DAYS)
    n = 0
    for job in ExportJob.objects.filter(status=ExportJob.Status.DONE, finished_at__lt=cutoff):
        delete_file(job)
        ExportJob.objects.filter(pk=job.pk).update(status=ExportJob.Status.EXPIRED, file_path="", file_size=0)
        n += 1
    return n


def delete_file(job):
    if job.file_path:
        try:
            (exports_dir() / job.file_path).unlink(missing_ok=True)
        except OSError:
            log.warning("Could not delete export file %s", job.file_path)


def run_pending():
    """Requeue abandoned jobs and run every pending one here (management command). Returns jobs run."""
    close_old_connections()
    requeue_stale()
    ran = 0
    for job_id in list(ExportJob.objects.filter(status=ExportJob.Status.PENDING).order_by("created_at")
                       .values_list("pk", flat=True)):
        if run_job(job_id) is not None:
            ran += 1
    return ran


def _describe_users(filters):
    from accounts.exports import _describe

    return _describe({"Search": filters.get("search", ""), "Role": filters.get("role", ""),
                      "Active": filters.get("is_active", "")})
