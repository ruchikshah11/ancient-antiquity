"""Pre-calculated daily stats (equal to the live queries, lazy refresh), FTS booking search, export jobs."""
import csv
import datetime
import io
import shutil
import tempfile
from decimal import Decimal
from unittest import mock

from django.core.management import call_command
from django.db import connection
from django.db.models import Sum
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook

from bookings import services
from bookings.models import Booking
from bookings.tests.factories import (
    APITestCase, at, book, client_for, make_admin, make_city, make_museum, make_staff, make_state, make_visitor,
)
from bookings.views import search_bookings
from stats import exports, legacy
from stats import services as stats
from stats.models import DailyStat, ExportJob, StatDay


def days_ago(n):
    return timezone.localdate() - datetime.timedelta(days=n)


def move(booking, created=None, **fields):
    """Change a booking behind the ORM's back (no signals), like old data or bulk updates."""
    if created is not None:
        fields["created_at"] = created
    Booking.objects.filter(pk=booking.pk).update(**fields)
    booking.refresh_from_db()
    return booking


class StatsData(APITestCase):
    """A spread of bookings: places in two states, channels, groups, statuses, refunds, check-ins, several days."""

    def setUp(self):
        super().setUp()
        self.museum_a = make_museum("Amber Fort", city=make_city(make_state("Rajasthan")))
        self.museum_b = make_museum("Bibi Ka Maqbara", city=make_city(make_state("Maharashtra")))
        self.staff = make_staff(self.museum_a)
        self.admin = make_admin()
        self.visitor = make_visitor()
        today = timezone.localdate()
        self.bookings = []
        specs = [
            # (museum, qty, confirm, days ago created, extra)
            (self.museum_a, 2, True, 0, {}),
            (self.museum_a, 3, True, 1, {"discount_amount": Decimal("10.00"), "addons_amount": Decimal("25.50")}),
            (self.museum_b, 4, True, 3, {"channel": Booking.Channel.COUNTER, "membership_discount": Decimal("7.25")}),
            (self.museum_b, 5, True, 10, {"is_group": True, "organisation_name": "St. Xavier's School"}),
            (self.museum_a, 1, False, 2, {}),                               # pending
            (self.museum_b, 6, True, 20, {"channel": Booking.Channel.AGENT}),
            (self.museum_a, 7, True, 40, {}),                               # outside 30 days
        ]
        for museum, qty, confirm, ago, extra in specs:
            b = book(self.visitor, museum, qty=qty, confirm=confirm)
            move(b, created=at(today - datetime.timedelta(days=ago), 11, 30), **extra)
            self.bookings.append(b)
        # cancelled with a refund, checked in today, a booking created just before midnight (time-zone edge)
        refunded = book(self.visitor, self.museum_b, qty=2, confirm=True)
        services.cancel_booking(refunded)
        move(refunded, created=at(today - datetime.timedelta(days=5), 9))
        arrived = book(self.visitor, self.museum_a, qty=3, confirm=True, visit_date=today)
        services.mark_checked_in(arrived, self.staff)
        late = book(self.visitor, self.museum_a, qty=1, confirm=True, visit_date=today)
        move(late, created=at(today - datetime.timedelta(days=4), 23, 59))
        self.bookings += [refunded, arrived, late]
        call_command("rebuild_stats", verbosity=0)


class StatsEqualLiveQueriesTests(StatsData):
    def report(self, user, first, last, museum=None):
        params = {"from": str(first), "to": str(last)}
        if museum:
            params["museum"] = museum.slug
        r = client_for(user).get("/api/admin/reports/summary/", params)
        self.assertEqual(r.status_code, 200, r.data)
        return r.data

    def assert_report_matches(self, user, first, last, museum=None):
        new = self.report(user, first, last, museum)
        old = legacy.report(user, first, last, museum)
        self.assertEqual(legacy.differences({k: new[k] for k in old}, old), [])
        return new

    def test_reports_match_for_several_ranges_places_and_users(self):
        data = self.assert_report_matches(self.admin, days_ago(29), days_ago(0))
        self.assertGreater(data["totals"]["bookings"], 3)
        self.assertGreater(data["totals"]["refunds"], 0)
        self.assertEqual(data["totals"]["cancelled"], 1)
        self.assert_report_matches(self.admin, days_ago(60), days_ago(0))
        self.assert_report_matches(self.admin, days_ago(5), days_ago(1))
        self.assert_report_matches(self.admin, days_ago(0), days_ago(0))
        self.assert_report_matches(self.admin, days_ago(60), days_ago(0), self.museum_b)
        self.assert_report_matches(self.staff, days_ago(60), days_ago(0))
        self.assert_report_matches(self.staff, days_ago(60), days_ago(0), self.museum_a)

    def test_dashboard_matches(self):
        for user in (self.admin, self.staff):
            r = client_for(user).get("/api/admin/dashboard/")
            self.assertEqual(r.status_code, 200, r.data)
            old = legacy.dashboard(user, timezone.localdate())
            self.assertEqual(legacy.differences({k: r.data[k] for k in old}, old), [], user)
        r = client_for(self.admin).get("/api/admin/dashboard/")
        self.assertEqual(r.data["totals"]["checked_in_today"], 3)
        self.assertEqual(r.data["totals"]["expected_today"], 4)
        self.assertEqual(r.data["status_counts"]["checked_in"], 1)

    def test_rebuild_range_and_row_shape(self):
        DailyStat.objects.all().delete()
        call_command("rebuild_stats", "--from", str(days_ago(3)), "--to", str(days_ago(0)), verbosity=0)
        self.assertEqual(set(DailyStat.objects.values_list("date", flat=True)) - {days_ago(n) for n in range(4)}, set())
        row = DailyStat.objects.get(date=days_ago(1), museum=self.museum_a, status_group="paid")
        self.assertEqual((row.bookings, row.discounts, row.addons), (1, Decimal("10.00"), Decimal("25.50")))


class LazyRefreshTests(StatsData):
    def total(self, first=None):
        r = client_for(self.admin).get("/api/admin/reports/summary/", {"from": str(first or days_ago(29))})
        return r.data["totals"]

    def test_new_booking_marks_its_days_dirty_and_is_counted_on_read(self):
        before = self.total()["bookings"]
        b = book(self.visitor, self.museum_b, qty=1, confirm=True)
        self.assertTrue(StatDay.objects.filter(date=timezone.localdate(), dirty_since__isnull=False).exists())
        self.assertTrue(StatDay.objects.filter(date=b.visit_date, dirty_since__isnull=False).exists())
        self.assertEqual(self.total()["bookings"], before + 1)
        self.assertFalse(StatDay.objects.filter(date=timezone.localdate(), dirty_since__isnull=False).exists())

    def test_old_day_change_through_save_is_refreshed(self):
        old = self.bookings[5]  # 20 days ago, paid
        before = self.total()
        services.cancel_booking(old)  # saves the booking: post_save marks its sale day dirty
        self.assertTrue(StatDay.objects.filter(date=days_ago(20), dirty_since__isnull=False).exists())
        after = self.total()
        self.assertEqual(after["bookings"], before["bookings"] - 1)
        self.assertEqual(after["cancelled"], before["cancelled"] + 1)

    def test_reschedule_refreshes_old_and_new_visit_day(self):
        b = book(self.visitor, self.museum_a, qty=2, confirm=True, visit_date=timezone.localdate())
        client_for(self.admin).get("/api/admin/dashboard/")
        b.visit_date = timezone.localdate() + datetime.timedelta(days=3)
        b.save()
        self.assertTrue(StatDay.objects.filter(date=timezone.localdate(), dirty_since__isnull=False).exists())
        stats.ensure_fresh(all_dirty=True)
        self.assertEqual(
            DailyStat.objects.filter(date=timezone.localdate()).aggregate(n=Sum("visit_visitors"))["n"],
            Booking.objects.filter(visit_date=timezone.localdate()).aggregate(n=Sum("total_visitors"))["n"],
        )

    def test_today_is_recomputed_when_older_than_two_minutes(self):
        b = self.bookings[0]  # created today, paid
        before = self.total()["revenue"]
        move(b, total_amount=b.total_amount + 1000)  # .update(): no signal
        self.assertEqual(self.total()["revenue"], before)  # still fresh: not recomputed yet
        StatDay.objects.filter(date=timezone.localdate()).update(refreshed_at=timezone.now() - datetime.timedelta(minutes=3))
        self.assertEqual(self.total()["revenue"], before + 1000)

    def test_timer_refresh_runs_in_background_outside_a_transaction(self):
        StatDay.objects.filter(date=timezone.localdate()).update(refreshed_at=timezone.now() - datetime.timedelta(minutes=3))
        with mock.patch("stats.services.refresh_in_background") as background,                 mock.patch("stats.services.compute_day") as compute,                 mock.patch.object(type(connection), "in_atomic_block", new=False, create=True):
            stats.ensure_fresh(days_ago(29), days_ago(0))
        background.assert_called_once_with({timezone.localdate()})
        compute.assert_not_called()

    def test_old_day_bulk_update_waits_for_rebuild(self):
        b = self.bookings[5]
        before = self.total()["revenue"]
        move(b, total_amount=b.total_amount + 500)
        self.assertEqual(self.total()["revenue"], before)
        call_command("rebuild_stats", "--from", str(days_ago(20)), "--to", str(days_ago(20)), verbosity=0)
        self.assertEqual(self.total()["revenue"], before + 500)

    def test_missing_days_in_a_report_range_are_computed(self):
        DailyStat.objects.all().delete()
        StatDay.objects.all().delete()
        data = client_for(self.admin).get("/api/admin/reports/summary/", {"from": str(days_ago(60))}).data
        old = legacy.report(self.admin, days_ago(60), days_ago(0))
        self.assertEqual(legacy.differences({k: data[k] for k in old}, old), [])

    def test_dashboard_marks_expired_holds(self):
        pending = self.bookings[4]
        move(pending, created=timezone.now() - datetime.timedelta(days=2))
        call_command("rebuild_stats", verbosity=0)
        client_for(self.admin).get("/api/admin/dashboard/")
        pending.refresh_from_db()
        self.assertEqual(pending.status, Booking.Status.EXPIRED)
        self.assertEqual(
            DailyStat.objects.filter(status_group="expired").aggregate(n=Sum("bookings"))["n"], 1,
        )


class SearchTests(APITestCase):
    def setUp(self):
        super().setUp()
        self.place = make_museum("Golconda Fort")
        self.other = make_museum("Red Fort")
        self.visitor = make_visitor()
        self.ravi = book(self.visitor, self.place, visitor_name="Ravi Kumar", visitor_phone="+91 98480 22338",
                         visitor_email="ravi.kumar@example.com")
        self.meena = book(self.visitor, self.other, visitor_name="Meena Iyer", visitor_phone="9000012345",
                          visitor_email="meena@example.org")
        self.group = book(self.visitor, self.other, visitor_name="Teacher", organisation_name="Delhi Public School",
                          visitor_phone="7012345678")

    def refs(self, term):
        return set(search_bookings(Booking.objects.all(), term).values_list("reference", flat=True))

    def test_uses_fts_on_sqlite(self):
        self.assertEqual(connection.vendor, "sqlite")
        with connection.cursor() as c:
            c.execute("SELECT count(*) FROM booking_search")
            self.assertEqual(c.fetchone()[0], 3)

    def test_name_organisation_place(self):
        self.assertEqual(self.refs("ravi"), {self.ravi.reference})
        self.assertEqual(self.refs("Ravi Kum"), {self.ravi.reference})
        self.assertEqual(self.refs("kumar"), {self.ravi.reference})
        self.assertEqual(self.refs("public school"), {self.group.reference})
        self.assertEqual(self.refs("golconda"), {self.ravi.reference})
        self.assertEqual(self.refs("red fort"), {self.meena.reference, self.group.reference})
        self.assertEqual(self.refs("nobody"), set())
        self.assertEqual(self.refs('"quoted" (odd) *'), set())  # no FTS syntax errors

    def test_phone_last_ten_digits_and_partials(self):
        self.assertEqual(self.refs("9848022338"), {self.ravi.reference})
        self.assertEqual(self.refs("+91 98480 22338"), {self.ravi.reference})
        self.assertEqual(self.refs("0919848022338"), {self.ravi.reference})
        self.assertEqual(self.refs("984802"), {self.ravi.reference})    # how it starts
        self.assertEqual(self.refs("022338"), {self.ravi.reference})    # how it ends
        self.assertEqual(self.refs("12345"), set())  # under 6 digits: treated as text
        self.assertEqual(self.refs("012345"), {self.meena.reference})

    def test_email_and_reference(self):
        self.assertEqual(self.refs("ravi.kumar@"), {self.ravi.reference})
        self.assertEqual(self.refs("RAVI.KUMAR@example.com"), {self.ravi.reference})
        self.assertEqual(self.refs("kumar@example.com"), set())  # email is a prefix match
        self.assertEqual(self.refs(self.meena.reference), {self.meena.reference})
        self.assertEqual(self.refs(self.meena.reference[:7]), {
            b.reference for b in (self.ravi, self.meena, self.group) if b.reference.startswith(self.meena.reference[:7])
        })
        self.assertEqual(self.refs(self.meena.reference[3:]), {self.meena.reference})  # without "AA-"

    def test_triggers_keep_index_in_sync(self):
        # update
        Booking.objects.filter(pk=self.meena.pk).update(visitor_name="Meenakshi Pillai", visitor_phone="8888877777")
        self.assertEqual(self.refs("pillai"), {self.meena.reference})
        self.assertEqual(self.refs("iyer"), set())
        self.assertEqual(self.refs("8888877777"), {self.meena.reference})
        self.assertEqual(self.refs("9000012345"), set())
        # a status-only save keeps the row
        services.confirm_payment(self.meena, Booking.PaymentMethod.UPI)
        self.assertEqual(self.refs("pillai"), {self.meena.reference})
        # place rename
        self.place.name = "Golkonda Qila"
        self.place.save()
        self.assertEqual(self.refs("qila"), {self.ravi.reference})
        self.assertEqual(self.refs("golconda"), set())
        # insert
        new = book(self.visitor, self.place, visitor_name="Zubin Mehta", visitor_phone="9123456780")
        self.assertEqual(self.refs("zubin"), {new.reference})
        # delete
        new.items.all().delete()
        new.addons.all().delete()
        Booking.objects.filter(pk=new.pk).delete()
        self.assertEqual(self.refs("zubin"), set())
        with connection.cursor() as c:
            c.execute("SELECT count(*) FROM booking_search")
            self.assertEqual(c.fetchone()[0], 3)

    def test_api_search(self):
        r = client_for(make_admin()).get("/api/bookings/", {"search": "meena"})
        self.assertEqual([b["reference"] for b in r.data["results"]], [self.meena.reference])


@override_settings(EXPORT_RUN_INLINE=True)
class ExportJobTests(APITestCase):
    url = "/api/admin/exports/"

    def setUp(self):
        super().setUp()
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        override = override_settings(EXPORTS_DIR=self.tmp)
        override.enable()
        self.addCleanup(override.disable)
        self.mine = make_museum("Assigned Place")
        self.other = make_museum("Other Place")
        self.admin = make_admin()
        self.staff = make_staff(self.mine)
        self.visitor = make_visitor()
        self.b1 = book(self.visitor, self.mine, qty=2, confirm=True, visitor_name="Kiran Das")
        self.b2 = book(self.visitor, self.other, qty=1, confirm=True)
        self.b3 = book(self.visitor, self.mine, qty=1)

    def start(self, user, kind, **filters):
        r = client_for(user).post(self.url, {"kind": kind, "filters": filters}, format="json")
        return r

    def download(self, user, job):
        return client_for(user).get(f"{self.url}{job['id']}/download/")

    def test_bookings_csv(self):
        r = self.start(self.admin, "bookings_csv")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual((r.data["status"], r.data["row_count"]), ("done", 3))
        resp = self.download(self.admin, r.data)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("attachment;", resp["Content-Disposition"])
        text = b"".join(resp.streaming_content).decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
        self.assertEqual(rows[0][0], "Reference")
        self.assertEqual({row[0] for row in rows[1:]}, {self.b1.reference, self.b2.reference, self.b3.reference})

    def test_bookings_xlsx_scoped_and_filtered(self):
        r = self.start(self.staff, "bookings_xlsx", search="kiran")
        self.assertEqual(r.data["status"], "done", r.data)
        wb = load_workbook(io.BytesIO(b"".join(self.download(self.staff, r.data).streaming_content)))
        refs = [row[0] for row in wb["Bookings"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(refs, [self.b1.reference])
        r = self.start(self.staff, "bookings_csv", museum=self.other.slug)
        self.assertEqual(r.status_code, 403)

    @override_settings(EXPORT_XLSX_SHEET_ROWS=2)
    def test_xlsx_continues_on_more_sheets(self):
        r = self.start(self.admin, "bookings_xlsx")
        wb = load_workbook(io.BytesIO(b"".join(self.download(self.admin, r.data).streaming_content)))
        self.assertEqual(wb.sheetnames, ["Bookings", "Bookings 2", "About"])
        self.assertIn("2 sheets", r.data["note"])

    @override_settings(EXPORT_PDF_MAX_ROWS=2)
    def test_pdf_is_capped_with_a_message(self):
        r = self.start(self.admin, "bookings_pdf")
        self.assertEqual((r.data["status"], r.data["row_count"], r.data["total_rows"]), ("done", 2, 3))
        self.assertIn("first 2 of 3", r.data["note"])
        self.assertTrue(b"".join(self.download(self.admin, r.data).streaming_content).startswith(b"%PDF"))

    def test_users_exports_admin_only(self):
        for kind in ("users_csv", "users_xlsx", "users_pdf"):
            r = self.start(self.admin, kind, role="staff")
            self.assertEqual(r.data["status"], "done", r.data)
            self.assertEqual(r.data["row_count"], 1)
        self.assertEqual(self.start(self.staff, "users_csv").status_code, 403)

    def test_download_needs_the_owner(self):
        job = self.start(self.admin, "bookings_csv").data
        self.assertEqual(self.download(make_admin(), job).status_code, 404)
        self.assertEqual(client_for().get(f"{self.url}{job['id']}/download/").status_code, 401)
        self.assertEqual(self.download(self.visitor, job).status_code, 403)
        listed = client_for(self.admin).get(self.url).data["results"]
        self.assertEqual([j["id"] for j in listed], [job["id"]])
        self.assertEqual(client_for(self.staff).get(self.url).data["results"], [])

    def test_bad_requests(self):
        self.assertEqual(self.start(self.admin, "bookings_doc").status_code, 400)
        self.assertEqual(self.start(self.admin, "bookings_csv", **{"from": "yesterday"}).status_code, 400)
        self.assertEqual(self.start(self.admin, "bookings_csv", museum="nowhere").status_code, 400)

    def test_files_are_deleted_after_seven_days(self):
        job = self.start(self.admin, "bookings_csv").data
        path = exports.exports_dir() / ExportJob.objects.get(pk=job["id"]).file_path
        self.assertTrue(path.exists())
        ExportJob.objects.filter(pk=job["id"]).update(finished_at=timezone.now() - datetime.timedelta(days=8))
        listed = client_for(self.admin).get(self.url).data["results"]
        self.assertEqual(listed[0]["status"], "expired")
        self.assertIsNone(listed[0]["download_url"])
        self.assertFalse(path.exists())
        self.assertEqual(self.download(self.admin, job).status_code, 404)

    def test_command_runs_pending_and_abandoned_jobs(self):
        pending = ExportJob.objects.create(user=self.admin, kind="bookings_csv")
        abandoned = ExportJob.objects.create(user=self.admin, kind="bookings_csv", status="running",
                                             heartbeat_at=timezone.now() - datetime.timedelta(hours=1))
        call_command("run_export_jobs", stdout=io.StringIO())
        for job in (pending, abandoned):
            job.refresh_from_db()
            self.assertEqual((job.status, job.row_count), ("done", 3))

    def test_background_thread_is_started_after_commit(self):
        with override_settings(EXPORT_RUN_INLINE=False), mock.patch("stats.exports.threading.Thread") as thread:
            with self.captureOnCommitCallbacks(execute=True):
                r = self.start(self.admin, "bookings_csv")
        self.assertEqual(r.data["status"], "pending")
        thread.assert_called_once()
        thread.return_value.start.assert_called_once()

    def test_failed_job_records_the_error(self):
        # A deliberate, fake failure; assertLogs keeps its traceback out of the test output.
        with mock.patch("stats.exports.write_bookings_csv", side_effect=RuntimeError("simulated export failure")),                 self.assertLogs("stats.exports", level="ERROR"):
            r = self.start(self.admin, "bookings_csv")
        self.assertEqual((r.data["status"], r.data["error"]), ("failed", "simulated export failure"))
