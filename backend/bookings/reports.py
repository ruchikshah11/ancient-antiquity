"""Panel reports (daily summary, CSV/Excel exports) and GST invoices.

Staff only see the places assigned to them; admins see everything.
"""
import csv
import datetime
import io
from decimal import ROUND_HALF_UP, Decimal

from django.conf import settings
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.permissions import IsPanelUser
from museums.models import Museum
from ancient_antiquity.dates import day_range

from .models import Booking
from .views import PAID, ensure_museum_access, panel_bookings

MAX_RANGE_DAYS = 366
CENT = Decimal("0.01")


def _money(value):
    return Decimal(value or 0).quantize(CENT, rounding=ROUND_HALF_UP)


def report_params(request):
    """Validated (from, to, museum) for ?from=&to=&museum=. Dates are booking (sale) dates."""
    params = request.query_params
    today = timezone.localdate()
    date_to = parse_date(params.get("to", "") or "") or today
    date_from = parse_date(params.get("from", "") or "") or date_to - datetime.timedelta(days=29)
    if date_from > date_to:
        raise ValidationError({"from": "The start date must be on or before the end date."})
    if (date_to - date_from).days >= MAX_RANGE_DAYS:
        raise ValidationError({"from": f"Choose a range of at most {MAX_RANGE_DAYS} days."})
    museum = None
    slug = (params.get("museum") or "").strip()
    if slug:
        museum = Museum.objects.filter(slug=slug).first()
        if not museum:
            raise ValidationError({"museum": "Unknown place."})
        ensure_museum_access(request.user, museum)
    return date_from, date_to, museum


def report_filters(request):
    """Validated (bookings queryset, from, to, museum) for ?from=&to=&museum=. Dates are booking (sale) dates."""
    date_from, date_to, museum = report_params(request)
    start, end = day_range(date_from, date_to)
    qs = panel_bookings(request.user).filter(created_at__gte=start, created_at__lt=end)
    if museum:
        qs = qs.filter(museum=museum)
    return qs, date_from, date_to, museum


class ReportSummaryView(APIView):
    """Report for a booking-date range, read from the pre-calculated stats.DailyStat (refreshed lazily)."""

    permission_classes = [IsPanelUser]

    def get(self, request):
        from stats import services as stats

        date_from, date_to, museum = report_params(request)
        rows = stats.panel_stats(request.user)
        if museum:
            rows = rows.filter(museum=museum)
        stats.ensure_fresh(date_from, date_to)
        by_day, splits, museum_rows = stats.report_summary(rows, date_from, date_to)

        daily = []
        t = dict(bookings=0, visitors=0, revenue=Decimal(0), refunds=Decimal(0), discounts=Decimal(0),
                 addons=Decimal(0), membership=Decimal(0), cancelled=0)
        day = date_from
        while day <= date_to:
            row = by_day.get(day, {})
            daily.append({
                "date": day,
                "bookings": row.get("bookings") or 0,
                "visitors": row.get("visitors") or 0,
                "revenue": _money(row.get("revenue")),
                "refunds": _money(row.get("refunds")),
            })
            for field in t:
                t[field] += row.get(field) or 0
            day += datetime.timedelta(days=1)

        totals = {
            "bookings": t["bookings"],
            "visitors": t["visitors"],
            "revenue": _money(t["revenue"]),
            "refunds": _money(t["refunds"]),
            "discounts": _money(t["discounts"]),
            # Revenue already includes add-ons (they are part of total_amount); these show the split.
            "addons": _money(t["addons"]),
            "membership_discounts": _money(t["membership"]),
            "cancelled": t["cancelled"],
        }
        totals["net"] = totals["revenue"]  # refunded bookings are cancelled, so already outside revenue

        def split(field, labels):
            rows = {}
            for r in splits:
                acc = rows.setdefault(r[field], {"bookings": 0, "visitors": 0, "revenue": Decimal(0)})
                acc["bookings"] += r["bookings"] or 0
                acc["visitors"] += r["visitors"] or 0
                acc["revenue"] += r["revenue"] or 0
            return [
                {
                    "key": key, "label": label,
                    "bookings": rows.get(key, {}).get("bookings", 0),
                    "visitors": rows.get(key, {}).get("visitors") or 0,
                    "revenue": _money(rows.get(key, {}).get("revenue")),
                }
                for key, label in labels
            ]

        info = stats.museum_info(r["museum_id"] for r in museum_rows)
        by_museum = sorted(
            (
                {
                    "slug": info[r["museum_id"]]["slug"], "name": info[r["museum_id"]]["name"],
                    "city": info[r["museum_id"]]["city__name"], "state": info[r["museum_id"]]["city__state__name"],
                    "bookings": r["bookings"], "visitors": r["visitors"] or 0, "revenue": _money(r["revenue"]),
                }
                for r in museum_rows if r["museum_id"] in info
            ),
            key=lambda m: (-m["revenue"], m["name"]),
        )
        return Response({
            "from": date_from,
            "to": date_to,
            "museum": {"slug": museum.slug, "name": museum.name} if museum else None,
            "totals": totals,
            "daily": daily,
            "by_channel": split("channel", Booking.Channel.choices),
            "by_type": split("is_group", [(False, "Individual"), (True, "Group")]),
            "by_museum": by_museum,
        })


EXPORT_HEADERS = [
    "Reference", "Created", "Visit date", "Slot", "Place", "City", "State", "Channel", "Group type",
    "Organisation", "Visitors", "Entered", "Amount", "Discount", "Payment method", "Payment ID",
    "Status", "Refund", "Refund status",
]


EXPORT_RELATED = ("museum__city__state", "slot", "event_session__event")


def export_rows(qs):
    qs = qs.select_related(*EXPORT_RELATED).order_by("created_at")
    for b in qs.iterator(chunk_size=500):
        yield booking_row(b)


def booking_row(b):
    """One export row (EXPORT_HEADERS order); booking needs EXPORT_RELATED selected."""
    return [
        b.reference,
        timezone.localtime(b.created_at).replace(tzinfo=None),
        b.visit_date,
        f"{b.slot_start_time:%H:%M}-{b.slot_end_time:%H:%M}" + (f" {b.event_title}" if b.event_session_id else ""),
        b.museum.name,
        b.museum.city.name,
        b.museum.city.state.name,
        b.get_channel_display(),
        b.get_group_type_display() if b.is_group else "",
        b.organisation_name,
        b.total_visitors,
        b.visitors_arrived,
        b.total_amount,
        b.discount_amount,
        b.get_payment_method_display() if b.payment_method else "",
        b.payment_id,
        b.get_status_display(),
        b.refund_amount,
        b.get_refund_status_display() if b.refund_status else "",
    ]


def _filename(date_from, date_to, museum, ext):
    part = f"-{museum.slug}" if museum else ""
    return f"bookings{part}-{date_from:%Y%m%d}-{date_to:%Y%m%d}.{ext}"


class BookingsCsvView(APIView):
    permission_classes = [IsPanelUser]

    def get(self, request):
        qs, date_from, date_to, museum = report_filters(request)
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{_filename(date_from, date_to, museum, "csv")}"'
        response.write("﻿")  # BOM so Excel opens the ₹/Hindi text as UTF-8
        writer = csv.writer(response)
        writer.writerow(EXPORT_HEADERS)
        for row in export_rows(qs):
            row[1] = row[1].strftime("%Y-%m-%d %H:%M")
            writer.writerow(row)
        return response


class BookingsXlsxView(APIView):
    permission_classes = [IsPanelUser]

    def get(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        qs, date_from, date_to, museum = report_filters(request)
        wb = Workbook()
        ws = wb.active
        ws.title = "Bookings"
        ws.append(EXPORT_HEADERS)
        header_fill = PatternFill("solid", fgColor="DCEFFB")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for row in export_rows(qs):
            ws.append(row)
        money_cols = [EXPORT_HEADERS.index(h) + 1 for h in ("Amount", "Discount", "Refund")]
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 2).number_format = "yyyy-mm-dd hh:mm"
            ws.cell(r, 3).number_format = "yyyy-mm-dd"
            for c in money_cols:
                ws.cell(r, c).number_format = "#,##0.00"
        widths = [13, 17, 12, 12, 34, 16, 18, 14, 18, 26, 9, 9, 11, 10, 15, 22, 13, 10, 14]
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        info = wb.create_sheet("About")
        info.append(["Ancient Antiquity bookings export"])
        info.append(["Booked from", date_from])
        info.append(["Booked to", date_to])
        info.append(["Place", museum.name if museum else "All places you can access"])
        info.append(["Generated", timezone.localtime().replace(tzinfo=None)])
        info["A1"].font = Font(bold=True, size=13)
        info.column_dimensions["A"].width = 16
        info.column_dimensions["B"].width = 32

        buffer = io.BytesIO()
        wb.save(buffer)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{_filename(date_from, date_to, museum, "xlsx")}"'
        return response


def financial_year(day):
    """Indian financial year label (April-March), e.g. 2026-27."""
    start = day.year if day.month >= 4 else day.year - 1
    return f"{start}-{str(start + 1)[-2:]}"


def build_invoice(booking):
    """Tax invoice data. The paid total is treated as GST-inclusive and split into taxable value + CGST/SGST."""
    paid_at = timezone.localtime(booking.paid_at or booking.created_at)
    gst_percent = Decimal(str(settings.INVOICE_GST_PERCENT))
    # Invoice value = what the visitor paid online + what a gift voucher covered (a voucher is prepayment).
    voucher = _money(getattr(booking, "voucher_amount", 0) or 0)
    group_discount = _money(booking.discount_amount)
    promo_discount = _money(getattr(booking, "promo_discount", 0) or 0)
    promo = getattr(booking, "promo_code", None)
    membership_discount = _money(getattr(booking, "membership_discount", 0) or 0)
    membership = getattr(booking, "membership", None)
    discount = group_discount + membership_discount + promo_discount
    total = _money(booking.total_amount + voucher)
    taxable = _money(total / (1 + gst_percent / 100))
    gst_total = total - taxable
    cgst = _money(gst_total / 2)
    sgst = gst_total - cgst
    half_rate = gst_percent / 2

    items = [
        {
            "description": item.ticket_name,
            "quantity": item.quantity,
            "unit_price": _money(item.unit_price),
            "amount": _money(item.unit_price * item.quantity),
            "kind": "ticket",
        }
        for item in booking.items.all()
        if item.quantity
    ]
    if booking.escort_count:
        items.append({
            "description": "Escort (free)", "quantity": booking.escort_count,
            "unit_price": _money(0), "amount": _money(0), "kind": "ticket",
        })
    # Add-ons (audio guide, camera permit...) are separate lines after the entry tickets.
    items += [
        {
            "description": line.name,
            "quantity": line.quantity,
            "unit_price": _money(line.unit_price),
            "amount": _money(line.unit_price * line.quantity),
            "kind": "addon",
        }
        for line in booking.addons.all()
    ]
    museum = booking.museum
    return {
        "invoice_number": f"AA/{financial_year(paid_at.date())}/{booking.reference}",
        "invoice_date": paid_at.date(),
        "reference": booking.reference,
        "seller": {
            "name": settings.INVOICE_SELLER_NAME,
            "address": settings.INVOICE_SELLER_ADDRESS,
            "gstin": settings.INVOICE_SELLER_GSTIN,
            "email": settings.SUPPORT_EMAIL,
        },
        "buyer": {
            "name": booking.visitor_name,
            "email": booking.visitor_email,
            "phone": booking.visitor_phone,
            "organisation": booking.organisation_name if booking.is_group else "",
            "group_type": booking.get_group_type_display() if booking.is_group else "",
        },
        "place": {
            "name": museum.name,
            "address": museum.address,
            "city": museum.city.name,
            "state": museum.city.state.name,
        },
        "place_of_supply": museum.city.state.name,
        "visit_date": booking.visit_date,
        "slot_start": booking.slot_start_time,
        "slot_end": booking.slot_end_time,
        "event_title": booking.event_title,
        "total_visitors": booking.total_visitors,
        "items": items,
        "subtotal": _money(total + discount),
        "discount": discount,
        # Separate lines so each saving is labelled correctly on the invoice.
        "group_discount": group_discount,
        "membership_discount": membership_discount,
        "membership_card": membership.card_number if membership else "",
        "promo_discount": promo_discount,
        "promo_code": promo.code if promo else "",
        "total": total,
        "voucher_amount": voucher,
        "paid_online": _money(booking.total_amount),
        "taxable_value": taxable,
        "gst_percent": gst_percent,
        "cgst_percent": half_rate,
        "cgst": cgst,
        "sgst_percent": half_rate,
        "sgst": sgst,
        "gst_total": gst_total,
        "payment": {
            "method": booking.get_payment_method_display() if booking.payment_method else "",
            "payment_id": booking.payment_id,
            "paid_at": booking.paid_at,
            "channel": booking.get_channel_display(),
        },
        "status": booking.status,
        "status_label": booking.get_status_display(),
        "notes": [
            "This is a computer-generated invoice and does not need a signature.",
            "Prices are inclusive of GST. The GST treatment shown here is indicative and must be confirmed with an accountant.",
        ],
    }
