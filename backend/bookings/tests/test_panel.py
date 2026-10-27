"""Panel scoping (staff vs admin), dashboard, reports/exports and GST invoices."""
import datetime
import io
from decimal import Decimal

from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook

from bookings import services
from bookings.models import Booking
from bookings.tests.factories import (
    ALL_DAY, APITestCase, book, client_for, make_admin, make_museum, make_staff, make_visitor,
)


class PanelBase(APITestCase):
    def setUp(self):
        super().setUp()
        self.mine = make_museum("Assigned Place")
        self.other = make_museum("Other Place")
        self.staff = make_staff(self.mine)
        self.admin = make_admin()
        self.visitor = make_visitor()
        self.paid_mine = book(self.visitor, self.mine, qty=2, confirm=True)          # 100
        self.paid_other = book(self.visitor, self.other, qty=4, confirm=True)        # 200
        self.pending_mine = book(self.visitor, self.mine, qty=1)


class StaffScopingTests(PanelBase):
    def refs(self, user, **params):
        r = client_for(user).get("/api/bookings/", params)
        self.assertEqual(r.status_code, 200, r.data)
        return {b["reference"] for b in r.data["results"]}

    def test_assigned_staff_sees_only_their_place(self):
        self.assertEqual(self.refs(self.staff), {self.paid_mine.reference, self.pending_mine.reference})
        r = client_for(self.staff).get(f"/api/bookings/{self.paid_other.reference}/")
        self.assertEqual(r.status_code, 404)

    def test_admin_sees_everything(self):
        self.assertEqual(len(self.refs(self.admin)), 3)

    def test_unassigned_staff_is_told_why(self):
        r = client_for(make_staff()).get("/api/bookings/")
        self.assertEqual(r.status_code, 403)
        self.assertIn("not assigned", str(r.data["detail"]))

    def test_mine_param_limits_panel_users_to_own_bookings(self):
        own = book(self.staff, self.other, confirm=True)
        self.assertEqual(self.refs(self.staff, mine=1), {own.reference})


class DashboardTests(PanelBase):
    url = "/api/admin/dashboard/"

    def test_admin_totals(self):
        r = client_for(self.admin).get(self.url)
        self.assertEqual(r.status_code, 200)
        totals = r.data["totals"]
        self.assertEqual(Decimal(totals["revenue"]), Decimal("300.00"))
        self.assertEqual(totals["bookings"], 2)
        self.assertEqual(totals["visitors"], 6)
        self.assertEqual(r.data["status_counts"], {"confirmed": 2, "pending": 1})
        self.assertEqual(len(r.data["daily"]), 14)

    def test_staff_totals_are_scoped(self):
        r = client_for(self.staff).get(self.url)
        self.assertEqual(Decimal(r.data["totals"]["revenue"]), Decimal("100.00"))
        self.assertEqual(r.data["totals"]["museums"], 1)

    def test_permissions(self):
        self.assertEqual(client_for(self.visitor).get(self.url).status_code, 403)
        self.assertEqual(client_for(make_staff()).get(self.url).status_code, 403)
        self.assertEqual(client_for().get(self.url).status_code, 401)


class ReportTests(PanelBase):
    def setUp(self):
        super().setUp()
        services.cancel_booking(self.paid_other)  # full refund of 200

    def test_summary(self):
        r = client_for(self.admin).get("/api/admin/reports/summary/")
        self.assertEqual(r.status_code, 200, r.data)
        t = r.data["totals"]
        self.assertEqual((t["bookings"], t["visitors"]), (1, 2))
        self.assertEqual(t["revenue"], Decimal("100.00"))
        self.assertEqual(t["refunds"], Decimal("200.00"))
        self.assertEqual(t["cancelled"], 1)
        self.assertEqual(len(r.data["daily"]), 30)
        online = next(c for c in r.data["by_channel"] if c["key"] == "online")
        self.assertEqual(online["bookings"], 1)

    def test_summary_scoped_for_staff(self):
        r = client_for(self.staff).get("/api/admin/reports/summary/")
        self.assertEqual(r.data["totals"]["refunds"], Decimal("0.00"))
        r = client_for(self.staff).get("/api/admin/reports/summary/", {"museum": self.other.slug})
        self.assertEqual(r.status_code, 403)

    def test_bad_ranges(self):
        today = timezone.localdate()
        r = client_for(self.admin).get("/api/admin/reports/summary/", {"from": str(today), "to": str(today - datetime.timedelta(days=1))})
        self.assertEqual(r.status_code, 400)
        r = client_for(self.admin).get("/api/admin/reports/summary/", {"from": str(today - datetime.timedelta(days=400))})
        self.assertEqual(r.status_code, 400)

    def test_csv_export(self):
        r = client_for(self.admin).get("/api/admin/reports/bookings.csv")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r["Content-Type"].startswith("text/csv"))
        self.assertIn("attachment;", r["Content-Disposition"])
        text = r.content.decode("utf-8")
        self.assertTrue(text.startswith("﻿Reference,"))
        self.assertIn(self.paid_mine.reference, text)
        self.assertEqual(len(text.strip().splitlines()), 4)  # header + 3 bookings

    def test_xlsx_export(self):
        r = client_for(self.staff).get("/api/admin/reports/bookings.xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb = load_workbook(io.BytesIO(r.content))
        rows = list(wb["Bookings"].iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "Reference")
        self.assertEqual({row[0] for row in rows[1:]}, {self.paid_mine.reference, self.pending_mine.reference})

    def test_exports_need_panel_user(self):
        self.assertEqual(client_for(self.visitor).get("/api/admin/reports/bookings.csv").status_code, 403)


@override_settings(INVOICE_GST_PERCENT=18.0)
class InvoiceTests(APITestCase):
    def setUp(self):
        super().setUp()
        self.museum = make_museum(tickets=[("Indian adult", "590.00"), ("Foreign national", "100.00")],
                                  slots=[(*ALL_DAY, 200)])
        self.visitor = make_visitor()

    def invoice(self, booking, user=None):
        return client_for(user or self.visitor).get(f"/api/bookings/{booking.reference}/invoice/")

    def test_only_paid_bookings(self):
        pending = book(self.visitor, self.museum)
        r = self.invoice(pending)
        self.assertEqual(r.status_code, 400)
        Booking.objects.filter(pk=pending.pk).update(status=Booking.Status.CANCELLED)
        self.assertEqual(self.invoice(pending).status_code, 400)

    def test_gst_split(self):
        booking = book(self.visitor, self.museum, qty=2, confirm=True)  # 1180 incl. 18% GST
        r = self.invoice(booking)
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["total"], Decimal("1180.00"))
        self.assertEqual(r.data["taxable_value"], Decimal("1000.00"))
        self.assertEqual((r.data["cgst"], r.data["sgst"]), (Decimal("90.00"), Decimal("90.00")))
        self.assertEqual(r.data["cgst_percent"], Decimal("9"))
        self.assertTrue(r.data["invoice_number"].startswith("AA/") and r.data["invoice_number"].endswith(booking.reference))

    def test_gst_split_with_rounding_adds_up(self):
        booking = book(self.visitor, self.museum, qty=1, ticket_name="Foreign national", confirm=True)  # 100
        d = self.invoice(booking).data
        self.assertEqual(d["taxable_value"], Decimal("84.75"))
        self.assertEqual((d["cgst"], d["sgst"]), (Decimal("7.63"), Decimal("7.62")))
        self.assertEqual(d["taxable_value"] + d["cgst"] + d["sgst"], d["total"])

    def test_group_invoice_lists_discount_and_free_escorts(self):
        booking = book(self.visitor, self.museum, qty=21, ticket_name="Foreign national", confirm=True,
                       group_type=Booking.GroupType.SCHOOL, organisation_name="DPS", escort_count=3)
        d = self.invoice(booking).data
        self.assertEqual(d["subtotal"], Decimal("2100.00"))
        self.assertEqual(d["discount"], Decimal("525.00"))
        self.assertEqual(d["total"], Decimal("1575.00"))
        self.assertEqual(d["items"][-1]["description"], "Escort (free)")
        self.assertEqual(d["buyer"]["organisation"], "DPS")

    def test_other_visitors_cannot_fetch_invoice(self):
        booking = book(self.visitor, self.museum, confirm=True)
        self.assertEqual(self.invoice(booking, make_visitor()).status_code, 404)
        self.assertEqual(self.invoice(booking, make_staff(self.museum)).status_code, 200)
